from pathlib import Path
from flask import Flask, render_template, request, redirect, url_for, flash
from openpyxl import Workbook, load_workbook

app = Flask(__name__)
app.secret_key = "cambia-esta-clave-en-produccion"

BASE_DIR = Path(__file__).resolve().parent
ARCHIVO_EXCEL = BASE_DIR / "data" / "estudiantes.xlsx"


def asegurar_excel():
    """Crea el archivo Excel y su hoja si todavía no existen."""
    ARCHIVO_EXCEL.parent.mkdir(parents=True, exist_ok=True)

    if not ARCHIVO_EXCEL.exists():
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Estudiantes"
        hoja.append(["ID", "Nombre", "Edad"])
        libro.save(ARCHIVO_EXCEL)
        return

    libro = load_workbook(ARCHIVO_EXCEL)
    if "Estudiantes" not in libro.sheetnames:
        hoja = libro.create_sheet("Estudiantes")
        hoja.append(["ID", "Nombre", "Edad"])
        libro.save(ARCHIVO_EXCEL)
    libro.close()


def leer_estudiantes():
    """Lee todos los estudiantes almacenados en Excel."""
    asegurar_excel()
    libro = load_workbook(ARCHIVO_EXCEL, read_only=True, data_only=True)
    hoja = libro["Estudiantes"]

    estudiantes = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0] is not None:
            estudiantes.append({
                "id": fila[0],
                "nombre": fila[1] or "",
                "edad": fila[2],
            })

    libro.close()
    return estudiantes


def siguiente_id():
    estudiantes = leer_estudiantes()
    ids = [int(e["id"]) for e in estudiantes if str(e["id"]).isdigit()]
    return max(ids, default=0) + 1


def agregar_estudiante(nombre, edad):
    """Agrega un estudiante al archivo Excel."""
    asegurar_excel()
    libro = load_workbook(ARCHIVO_EXCEL)
    hoja = libro["Estudiantes"]
    hoja.append([siguiente_id(), nombre, edad])
    libro.save(ARCHIVO_EXCEL)
    libro.close()


def eliminar_estudiante(estudiante_id):
    """Elimina un estudiante por ID."""
    asegurar_excel()
    libro = load_workbook(ARCHIVO_EXCEL)
    hoja = libro["Estudiantes"]

    eliminado = False
    for fila in range(2, hoja.max_row + 1):
        if hoja.cell(fila, 1).value == estudiante_id:
            hoja.delete_rows(fila, 1)
            eliminado = True
            break

    libro.save(ARCHIVO_EXCEL)
    libro.close()
    return eliminado


@app.route("/")
def index():
    estudiantes = leer_estudiantes()
    return render_template(
        "index.html",
        estudiantes=estudiantes,
        total=len(estudiantes)
    )


@app.post("/registrar")
def registrar():
    nombre = request.form.get("nombre", "").strip()
    edad_texto = request.form.get("edad", "").strip()

    if not nombre:
        flash("El nombre del estudiante es obligatorio.", "error")
        return redirect(url_for("index"))

    if not edad_texto.isdigit():
        flash("La edad debe ser un número entero válido.", "error")
        return redirect(url_for("index"))

    edad = int(edad_texto)

    if edad < 1 or edad > 120:
        flash("Ingrese una edad entre 1 y 120 años.", "error")
        return redirect(url_for("index"))

    agregar_estudiante(nombre, edad)
    flash(f"Estudiante {nombre} registrado correctamente.", "success")
    return redirect(url_for("index"))


@app.post("/eliminar/<int:estudiante_id>")
def eliminar(estudiante_id):
    if eliminar_estudiante(estudiante_id):
        flash("Estudiante eliminado correctamente.", "success")
    else:
        flash("No se encontró el estudiante.", "error")
    return redirect(url_for("index"))


@app.get("/salud")
def salud():
    return {"status": "ok", "servicio": "registro-estudiantes"}


if __name__ == "__main__":
    asegurar_excel()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
